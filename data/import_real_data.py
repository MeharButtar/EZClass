"""
Turns the team's filled-in template (one easy sheet) into the CSV files the
database loads. This is the bridge between "teammates typing into a
spreadsheet" and the six-table schema, so nobody has to deal with ids.

Add --skip-problems to import only the rows that pass and list the rest.

Usage (from the project root):
    python3 data/import_real_data.py data/template/EzClass_Data_Template.xlsx
    python3 data/import_real_data.py my_rows.csv          # a CSV also works

It checks every row, prints any problems with the row number, and only
writes files if everything is valid. Output goes to data/cleaned/:
    professors.csv, rooms.csv, courses.csv, sections.csv

Then load them with:
    python3 backend/models/db.py data/cleaned
"""

import csv
import os
import sys

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, BASE)
from backend.models.conflicts import parse_days  # noqa: E402

OUT_DIR = os.path.join(BASE, "data", "cleaned")

COLUMNS = ["term", "subject", "catalog_num", "course_title", "units",
           "section_number", "days", "start_time", "end_time",
           "professor_first", "professor_last", "building", "room_number",
           "capacity"]
# Headers in the template are friendly ("Course Number (e.g. 150)"). Only the
# part before the first " (" is used to find the column, so the hint text can
# change. The old code-style names still work too.
FRIENDLY = {
    "term": "term", "subject": "subject", "course number": "catalog_num",
    "course title": "course_title", "units": "units", "section": "section_number",
    "days": "days", "start time": "start_time", "end time": "end_time",
    "professor first name": "professor_first", "professor last name": "professor_last",
    "building": "building", "room number": "room_number", "seats": "capacity",
}


def header_key(text):
    base = text.split(" (")[0].strip().lower()
    return FRIENDLY.get(base, text.strip())


REQUIRED = ["term", "subject", "catalog_num", "course_title", "units",
            "section_number", "days", "start_time", "end_time"]


def read_rows(path):
    """Returns a list of dicts keyed by COLUMNS. Works for .xlsx and .csv."""
    if path.lower().endswith(".xlsx"):
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb["Sections"]
        header = [header_key(str(c.value)) if c.value is not None else "" for c in ws[1]]
        raw = [[c.value for c in row] for row in ws.iter_rows(min_row=2)]
    else:
        with open(path, newline="", encoding="utf-8-sig") as f:
            reader = csv.reader(f)
            header = [header_key(h) for h in next(reader)]
            raw = list(reader)
    missing = [c for c in COLUMNS if c not in header]
    if missing:
        sys.exit(f"Missing column(s) in the sheet header: {', '.join(missing)}")
    rows = []
    for i, values in enumerate(raw, start=2):
        rec = {h: ("" if v is None else str(v).strip()) for h, v in zip(header, values)}
        if not any(rec.get(c) for c in COLUMNS):
            continue  # blank row
        if rec.get("term", "").upper().startswith("EXAMPLE"):
            continue  # the gray sample row in the template
        rec["_row"] = i
        rows.append(rec)
    return rows


def norm_time(value):
    """'9:00', '09:00', '9:00:00' or an Excel time -> 'HH:MM', or None if bad."""
    v = value.strip()
    if " " in v:                      # '1899-12-30 09:00:00' style
        v = v.split(" ")[-1]
    parts = v.split(":")
    try:
        h, m = int(parts[0]), int(parts[1])
    except (ValueError, IndexError):
        return None
    if not (0 <= h <= 23 and 0 <= m <= 59):
        return None
    return f"{h:02d}:{m:02d}"


def strip_float(value):
    """Excel sometimes hands us 101.0 or 2.0 for numbers typed in a text column."""
    v = value.strip()
    if v.replace(".", "", 1).isdigit() and "." in v and float(v) == int(float(v)):
        return str(int(float(v)))
    return v


def clean_number(value):
    """'3.0' -> '3'"""
    try:
        f = float(value)
    except ValueError:
        return None
    return str(int(f)) if f == int(f) else None


def validate(rows):
    errors, good = [], []
    for r in rows:
        problems = []
        for col in REQUIRED:
            if not r.get(col):
                problems.append(f"{col} is empty")
        if problems:
            errors.append((r["_row"], "; ".join(problems)))
            continue

        r["subject"] = r["subject"].upper().replace(" ", "")
        r["catalog_num"] = strip_float(r["catalog_num"]).upper().replace(" ", "")
        r["days"] = r["days"].replace(" ", "")
        r["units"] = clean_number(r["units"])
        r["start_time"] = norm_time(r["start_time"])
        r["end_time"] = norm_time(r["end_time"])
        r["capacity"] = clean_number(r["capacity"]) if r.get("capacity") else ""
        r["room_number"] = strip_float(r.get("room_number", ""))   # Excel gives 22.0 for room 22
        r["section_number"] = strip_float(r["section_number"])
        if r["section_number"].isdigit():
            r["section_number"] = r["section_number"].zfill(2)
        # a building with no room (or the reverse) gets "TBA" for the missing half
        if r.get("building") and not r.get("room_number"):
            r["room_number"] = "TBA"
        elif r.get("room_number") and not r.get("building"):
            r["building"] = "TBA"

        if r["units"] is None:
            problems.append("units must be a whole number")
        if not parse_days(r["days"]):
            problems.append(f"days '{r['days']}' not understood (use MW, TTh, MWF, TR, M, F ...)")
        if r["start_time"] is None or r["end_time"] is None:
            problems.append("start_time/end_time must look like 14:30")
        elif r["start_time"] >= r["end_time"]:
            problems.append("end_time must be after start_time")
        if r.get("capacity") is None:
            problems.append("capacity must be a whole number or blank")
        if bool(r["professor_first"]) != bool(r["professor_last"]):
            problems.append("give both professor_first and professor_last, or leave both blank")

        if problems:
            errors.append((r["_row"], "; ".join(problems)))
        else:
            good.append(r)

    # same course + section + term twice is almost certainly a typo
    seen, unique = {}, []
    for r in good:
        key = (r["term"], r["subject"], r["catalog_num"], r["section_number"])
        if key in seen:
            errors.append((r["_row"], f"duplicate of row {seen[key]} ({key[1]} {key[2]} sec {key[3]})"))
            continue
        seen[key] = r["_row"]
        unique.append(r)
    return unique, errors


def build_tables(rows):
    professors, rooms, courses, sections = {}, {}, {}, []
    for r in rows:
        prof_id = room_id = ""
        if r["professor_first"]:
            key = (r["professor_first"], r["professor_last"])
            if key not in professors:
                professors[key] = len(professors) + 1
            prof_id = professors[key]
        if r["building"]:
            key = (r["building"], r["room_number"])
            if key not in rooms:
                rooms[key] = len(rooms) + 1
            room_id = rooms[key]
        ckey = (r["subject"], r["catalog_num"])
        if ckey not in courses:
            courses[ckey] = (len(courses) + 1, r["course_title"], r["units"])
        sections.append({
            "section_id": len(sections) + 1,
            "course_id": courses[ckey][0],
            "professor_id": prof_id,
            "room_id": room_id,
            "section_number": r["section_number"],
            "term": r["term"],
            "days": r["days"],
            "start_time": r["start_time"],
            "end_time": r["end_time"],
            "capacity": r["capacity"],
        })
    return professors, rooms, courses, sections


def write(path, fields, rows):
    with open(path, "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        w.writerows(rows)
    print(f"  wrote {len(rows):>4} rows -> {os.path.relpath(path, BASE)}")


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    skip = "--skip-problems" in sys.argv
    if len(args) != 1:
        sys.exit(__doc__)
    rows = read_rows(args[0])
    good, errors = validate(rows)
    print(f"Read {len(rows)} rows: {len(good)} ok, {len(errors)} with problems.")
    if errors:
        if skip:
            print("\nSkipping these rows (they are NOT imported):")
        else:
            print("\nFix these rows in the sheet (row numbers match the sheet), then run again.")
            print("Or add --skip-problems to import only the good rows.")
        for row, msg in sorted(errors):
            print(f"  row {row}: {msg}")
        if not skip:
            sys.exit(1)
        print(f"\nImporting the {len(good)} good rows.")

    professors, rooms, courses, sections = build_tables(good)
    os.makedirs(OUT_DIR, exist_ok=True)
    write(os.path.join(OUT_DIR, "professors.csv"),
          ["professor_id", "first_name", "last_name", "email", "department"],
          [{"professor_id": i, "first_name": k[0], "last_name": k[1], "email": "", "department": "Unknown"}
           for k, i in professors.items()])
    write(os.path.join(OUT_DIR, "rooms.csv"),
          ["room_id", "building", "room_number", "capacity"],
          [{"room_id": i, "building": k[0], "room_number": k[1], "capacity": ""}
           for k, i in rooms.items()])
    write(os.path.join(OUT_DIR, "courses.csv"),
          ["course_id", "subject", "catalog_num", "title", "units", "description"],
          [{"course_id": v[0], "subject": k[0], "catalog_num": k[1], "title": v[1],
            "units": v[2], "description": ""} for k, v in courses.items()])
    write(os.path.join(OUT_DIR, "sections.csv"),
          ["section_id", "course_id", "professor_id", "room_id", "section_number",
           "term", "days", "start_time", "end_time", "capacity"], sections)
    print("\nNext: python3 backend/models/db.py data/cleaned")


if __name__ == "__main__":
    main()
